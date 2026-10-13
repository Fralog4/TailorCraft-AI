import io
from typing import Optional
import pdfplumber
import docx

def parse_cv_file(file_bytes: bytes, filename: str) -> str:
    """
    Extracts plain text from PDF, DOCX or TXT files.
    """
    ext = filename.lower().split(".")[-1] if "." in filename else ""

    if ext == "pdf":
        return _parse_pdf(file_bytes)
    elif ext in ["docx", "doc"]:
        return _parse_docx(file_bytes)
    elif ext in ["txt", "md"]:
        return file_bytes.decode("utf-8", errors="replace")
    else:
        # Attempt fallback to PDF first, then UTF-8 text
        try:
            return _parse_pdf(file_bytes)
        except Exception:
            return file_bytes.decode("utf-8", errors="replace")

def _parse_pdf(file_bytes: bytes) -> str:
    text_parts = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page_idx, page in enumerate(pdf.pages):
            text = page.extract_text()
            if text:
                text_parts.append(text)
    
    extracted = "\n\n".join(text_parts).strip()
    if not extracted:
        raise ValueError("Could not extract any text from the provided PDF. It might be scanned or image-only.")
    return extracted

def _parse_docx(file_bytes: bytes) -> str:
    doc = docx.Document(io.BytesIO(file_bytes))
    text_parts = []
    
    # Extract paragraphs
    for p in doc.paragraphs:
        if p.text.strip():
            text_parts.append(p.text.strip())
            
    # Extract text from tables if any
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                text_parts.append(row_text)
                
    extracted = "\n".join(text_parts).strip()
    if not extracted:
        raise ValueError("Could not extract text from the provided DOCX document.")
    return extracted
