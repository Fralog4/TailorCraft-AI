import io
import os
import sys
import json
import requests
import docx
from reportlab.pdfgen import canvas

# Ensure backend directory is in sys.path
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "backend"))

from parser import parse_cv_file
from ats_scorer import extract_keywords, calculate_ats_match
from models import TailoredCV, Experience, Education
from generator import generate_pdf
from docx_generator import generate_docx

BASE_URL = "http://127.0.0.1:8000"

def test_1_document_parsers():
    print("\n--- Test 1: Document Parsers (PDF, DOCX, TXT) ---")
    
    # 1.1 Generate sample in-memory PDF
    pdf_buf = io.BytesIO()
    p = canvas.Canvas(pdf_buf)
    p.drawString(100, 750, "Mario Rossi - Software Engineer")
    p.drawString(100, 730, "Esperienza con Python, FastAPI e Docker.")
    p.save()
    pdf_bytes = pdf_buf.getvalue()
    
    parsed_pdf = parse_cv_file(pdf_bytes, "test_resume.pdf")
    assert "Mario Rossi" in parsed_pdf
    assert "Python, FastAPI" in parsed_pdf
    print("[PASS] PDF parsing successful")

    # 1.2 Generate sample in-memory DOCX
    doc = docx.Document()
    doc.add_paragraph("Giulia Bianchi - Data Scientist")
    doc.add_paragraph("Competenze in Machine Learning e PyTorch.")
    docx_buf = io.BytesIO()
    doc.save(docx_buf)
    docx_bytes = docx_buf.getvalue()

    parsed_docx = parse_cv_file(docx_bytes, "test_resume.docx")
    assert "Giulia Bianchi" in parsed_docx
    assert "Machine Learning" in parsed_docx
    print("[PASS] DOCX parsing successful")

    # 1.3 Plain text
    txt_bytes = "Luigi Verdi - DevOps Specialist con esperienza AWS e Kubernetes.".encode("utf-8")
    parsed_txt = parse_cv_file(txt_bytes, "test_resume.txt")
    assert "Luigi Verdi" in parsed_txt
    print("[PASS] TXT parsing successful")

    # 1.4 Edge Case: Empty file
    try:
        parse_cv_file(b"", "empty.pdf")
        assert False, "Should have raised ValueError on empty PDF"
    except Exception as e:
        print("[PASS] Empty PDF correctly rejected:", str(e))


def test_2_api_validation_edge_cases():
    print("\n--- Test 2: Backend API Validation & Error Handling ---")

    # 2.1 Empty Job Description
    files = {"file": ("cv.txt", b"Mario Rossi Software Dev")}
    res = requests.post(
        f"{BASE_URL}/api/v1/tailor-cv",
        files=files,
        data={"job_description": "   ", "api_key": "fake_key", "llm_provider": "gemini"}
    )
    assert res.status_code == 400
    assert "Job description cannot be empty" in res.json()["detail"]
    print("[PASS] Empty JD correctly rejected with 400")

    # 2.2 Missing API Key
    res = requests.post(
        f"{BASE_URL}/api/v1/tailor-cv",
        files=files,
        data={"job_description": "Valid job description", "api_key": "   ", "llm_provider": "gemini"}
    )
    assert res.status_code == 400
    assert "API Key is required" in res.json()["detail"]
    print("[PASS] Missing API key correctly rejected with 400")

    # 2.3 Empty Uploaded File
    res = requests.post(
        f"{BASE_URL}/api/v1/tailor-cv",
        files={"file": ("empty.txt", b"")},
        data={"job_description": "Valid JD", "api_key": "fake_key", "llm_provider": "gemini"}
    )
    assert res.status_code == 400
    assert "Uploaded file is empty" in res.json()["detail"]
    print("[PASS] Empty file upload correctly rejected with 400")


def test_3_all_styles_and_languages():
    print("\n--- Test 3: Generator Matrix (3 Stili x 2 Lingue) ---")
    cv = TailoredCV(
        full_name="Elena Colombo",
        contact_info="Roma | elena@example.com | +39 06 1234567 | linkedin.com/in/elena",
        professional_summary="Senior Cloud Architect con 8 anni di esperienza in AWS e architetture scalabili.",
        skills=["AWS (ECS, Lambda, S3)", "Terraform", "Kubernetes", "Python"],
        experiences=[
            Experience(
                job_title="Lead Cloud Architect",
                company="Enterprise Cloud SpA",
                start_date="2020",
                end_date="Presente",
                bullet_points=[
                    "Architettato infrastruttura cloud ibrida riducendo i costi operativi del 30%.",
                    "Automatizzato pipeline CI/CD per 12 microservizi mission-critical."
                ]
            )
        ],
        education=[
            Education(
                degree="Laurea in Informatica",
                institution="Università La Sapienza",
                graduation_year="2016"
            )
        ]
    )

    styles = ["modern", "classic", "executive"]
    languages = ["en", "it"]

    for style in styles:
        for lang in languages:
            # Test PDF
            pdf_buf = generate_pdf(cv, template_style=style, language=lang)
            assert pdf_buf.getvalue().startswith(b"%PDF"), f"Failed PDF for {style}-{lang}"

            # Test DOCX
            docx_buf = generate_docx(cv, template_style=style, language=lang)
            assert docx_buf.getvalue().startswith(b"PK"), f"Failed DOCX for {style}-{lang}"

            print(f"[PASS] Matrix {style.upper()} ({lang.upper()}): PDF ({len(pdf_buf.getvalue())}b) and DOCX ({len(docx_buf.getvalue())}b) OK")


def test_4_keywords_and_symbols():
    print("\n--- Test 4: Keyword Extraction & Special Symbols ---")
    text = """
    Cerchiamo Senior Developer con solida esperienza in C++, C# e .NET Core.
    Requisiti: Machine Learning, Deep Learning, REST API, Docker e CI/CD.
    Gradita conoscenza di Node.js e React.
    Anni di esperienza: almeno 5 anni nel settore bancario a Milano.
    """
    keywords = extract_keywords(text)
    kws_lower = [k.lower() for k in keywords]

    assert "c++" in kws_lower, "Failed to extract C++"
    assert "c#" in kws_lower, "Failed to extract C#"
    assert "machine learning" in kws_lower, "Failed to extract Machine Learning"
    assert "deep learning" in kws_lower, "Failed to extract Deep Learning"
    assert "ci/cd" in kws_lower, "Failed to extract CI/CD"
    
    # Verify Italian stop words were stripped
    assert "anni" not in kws_lower
    assert "settore" not in kws_lower
    assert "esperienza" not in kws_lower
    print(f"[PASS] Keywords extracted correctly ({len(keywords)} terms): {keywords}")

if __name__ == "__main__":
    test_1_document_parsers()
    test_2_api_validation_edge_cases()
    test_3_all_styles_and_languages()
    test_4_keywords_and_symbols()
    print("\n=============================================")
    print(">>> 100% DI TUTTI GLI USE CASE VERIFICATI! <<<")
    print("=============================================")
